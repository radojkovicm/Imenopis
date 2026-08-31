"""Newborn XLSX loader (PROJECT.md §11.1, Step 1 of §10 build order).

Parses one year's `najcesca-imena-dece-...xlsx` into normalized rows and
writes `data/normalized/newborn_<year>.csv` (';' separator, UTF-8 BOM, per the
project-wide CSV convention in §11.1).

Structure confirmed by Phase 0 (docs/DATA_NOTES.md §2): one sheet, 32 groups
of exactly 10 rows each, rank is positional (row 1 in group = rank 1 ... row
10 = rank 10), column A repeats the group label on every row, column B = girl
name, column C = boy name. No count column exists in any of the 5 years
(BRANCH A resolved: rank only) - the loader never invents one.

A cell containing whitespace (a name+surname typo, found once in the 2021
file - docs/DATA_NOTES.md §2.2) is stored verbatim, never silently split or
corrected - the parser never fills gaps or guesses (PROJECT.md §11.2).
"""

import csv
from dataclasses import dataclass
from pathlib import Path

import openpyxl

from src.ingest.geography import classify_label

EXPECTED_GROUP_SIZE = 10


@dataclass
class NewbornRow:
    year: int
    gender: str  # 'M' or 'F'
    source_form: str
    rank: int
    district_code: str | None  # None = Republic-level row
    flagged: bool  # True if source_form looks malformed (e.g. contains whitespace)


class XlsxStructureError(Exception):
    """Raised when a source file doesn't match the structure Phase 0 verified."""


def _iter_groups(rows: list[tuple]) -> list[tuple[str, list[tuple]]]:
    """Group consecutive rows by column-A label, skipping the title/header rows."""
    groups: list[tuple[str, list[tuple]]] = []
    current_label: str | None = None
    current_rows: list[tuple] = []

    for row in rows:
        a = row[0]
        b = row[1] if len(row) > 1 else None
        c = row[2] if len(row) > 2 else None
        if a:
            if current_label is not None and a != current_label:
                groups.append((current_label, current_rows))
                current_rows = []
            current_label = a
            current_rows.append((b, c))
        else:
            if current_label is not None and (b or c):
                current_rows.append((b, c))
            elif current_label is not None:
                groups.append((current_label, current_rows))
                current_label = None
                current_rows = []
    if current_label:
        groups.append((current_label, current_rows))

    # Drop the title pseudo-group (its label contains the publication title).
    return [g for g in groups if "Најчешћа" not in g[0]]


def parse_xlsx(path: Path, year: int) -> list[NewbornRow]:
    wb = openpyxl.load_workbook(path, data_only=True)
    if len(wb.sheetnames) != 1:
        raise XlsxStructureError(
            f"{path.name}: expected exactly 1 sheet, found {wb.sheetnames}"
        )
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True))
    groups = _iter_groups(rows)

    out: list[NewbornRow] = []
    for label, group_rows in groups:
        if len(group_rows) != EXPECTED_GROUP_SIZE:
            raise XlsxStructureError(
                f"{path.name}: group {label!r} has {len(group_rows)} rows, "
                f"expected {EXPECTED_GROUP_SIZE}"
            )
        kind, district_code = classify_label(label)
        if kind == "rollup":
            continue  # macro-region / NUTS-2 rollups are not persisted (see geography.py)

        for rank, (girl, boy) in enumerate(group_rows, start=1):
            for gender, name in (("F", girl), ("M", boy)):
                if name is None or str(name).strip() == "":
                    continue  # missing cell -> no row, never inferred
                name = str(name)
                flagged = " " in name.strip()
                out.append(
                    NewbornRow(
                        year=year,
                        gender=gender,
                        source_form=name,
                        rank=rank,
                        district_code=district_code,
                        flagged=flagged,
                    )
                )
    return out


def write_normalized_csv(rows: list[NewbornRow], out_path: Path) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with out_path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow(["year", "gender", "source_form", "rank", "district_code", "flagged"])
        for r in rows:
            writer.writerow(
                [r.year, r.gender, r.source_form, r.rank, r.district_code or "", "1" if r.flagged else "0"]
            )


def load_year(raw_path: Path, year: int, normalized_dir: Path) -> list[NewbornRow]:
    rows = parse_xlsx(raw_path, year)
    write_normalized_csv(rows, normalized_dir / f"newborn_{year}.csv")
    return rows
